import os
import tkinter as tk
from tkinter import messagebox, ttk
import openpyxl

EXCEL_FILE = "student_results.xlsx"

# ----------------- EXCEL FILE INITIALIZATION -----------------
def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Results"
        headers = [
            "Name", "Roll No.", "Class", "Subject 1", "Subject 2",
            "Subject 3", "Subject 4", "Subject 5", "Total Marks",
            "Percentage", "Result"
        ]
        ws.append(headers)
        wb.save(EXCEL_FILE)

# ----------------- LOGIC FUNCTIONS -----------------
def save_student():
    name = entry_name.get().strip()
    roll = entry_roll.get().strip()
    cls = entry_class.get().strip()
    
    if not name or not roll or not cls:
        messagebox.showerror("Error", "Please fill Name, Roll No., and Class!")
        return
        
    try:
        s1 = float(entry_s1.get())
        s2 = float(entry_s2.get())
        s3 = float(entry_s3.get())
        s4 = float(entry_s4.get())
        s5 = float(entry_s5.get())
    except ValueError:
        messagebox.showerror("Error", "Marks must be numeric values!")
        return

    # Check for Duplicate Roll No.
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Results"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[1]) == roll:
            messagebox.showerror("Error", f"Roll No. {roll} already exists!")
            return

    # Calculations
    total = s1 + s2 + s3 + s4 + s5
    percentage = round((total / 500) * 100, 2)
    result = "Pass" if (s1 >= 35 and s2 >= 35 and s3 >= 35 and s4 >= 35 and s5 >= 35) else "Fail"

    # Append to Excel
    ws.append([name, roll, cls, s1, s2, s3, s4, s5, total, f"{percentage}%", result])
    wb.save(EXCEL_FILE)
    
    messagebox.showinfo("Success", "Student Record Saved Successfully!")
    clear_add_inputs()

def clear_add_inputs():
    entry_name.delete(0, tk.END)
    entry_roll.delete(0, tk.END)
    entry_class.delete(0, tk.END)
    entry_s1.delete(0, tk.END)
    entry_s2.delete(0, tk.END)
    entry_s3.delete(0, tk.END)
    entry_s4.delete(0, tk.END)
    entry_s5.delete(0, tk.END)

def get_result():
    search_roll = entry_search_roll.get().strip()
    if not search_roll:
        messagebox.showerror("Error", "Please enter a Roll No. to search!")
        return

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Results"]
    
    found = False
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[1]) == search_roll:
            lbl_res_name.config(text=f"Name: {row[0]}")
            lbl_res_class.config(text=f"Class: {row[2]}")
            lbl_res_total.config(text=f"Total: {row[8]}")
            lbl_res_per.config(text=f"Percentage: {row[9]}")
            lbl_res_status.config(text=f"Result: {row[10]}")
            found = True
            break

    if not found:
        messagebox.showwarning("Not Found", f"❌ Student record with Roll No. {search_roll} not found.")

def show_all_results():
    for item in tree.get_children():
        tree.delete(item)
        
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Results"]
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        # Display: Name, Roll No., Class, Total, Percentage, Result
        tree.insert("", tk.END, values=(row[0], row[1], row[2], row[8], row[9], row[10]))

# ----------------- GUI SETUP -----------------
init_excel()

root = tk.Tk()
root.title("Student Result Management System")
root.geometry("700x550")

notebook = ttk.Notebook(root)
notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

# --- TAB 1: ADD STUDENT ---
tab1 = ttk.Frame(notebook)
notebook.add(tab1, text="1. Add Student")

labels = ["Name:", "Roll No.:", "Class:", "Subject 1:", "Subject 2:", "Subject 3:", "Subject 4:", "Subject 5:"]
entries = []

for i, label in enumerate(labels):
    ttk.Label(tab1, text=label).grid(row=i, column=0, padx=10, pady=5, sticky=tk.W)
    ent = ttk.Entry(tab1)
    ent.grid(row=i, column=1, padx=10, pady=5)
    entries.append(ent)

(entry_name, entry_roll, entry_class, entry_s1, entry_s2, entry_s3, entry_s4, entry_s5) = entries

ttk.Button(tab1, text="💾 Save", command=save_student).grid(row=8, column=0, columnspan=2, pady=15)

# --- TAB 2: GET RESULT ---
tab2 = ttk.Frame(notebook)
notebook.add(tab2, text="2. Get Result")

ttk.Label(tab2, text="Enter Roll No.:").grid(row=0, column=0, padx=10, pady=10)
entry_search_roll = ttk.Entry(tab2)
entry_search_roll.grid(row=0, column=1, padx=10, pady=10)

ttk.Button(tab2, text="🔍 Get Result", command=get_result).grid(row=0, column=2, padx=10, pady=10)

lbl_res_name = ttk.Label(tab2, text="Name: -", font=("Arial", 11, "bold"))
lbl_res_name.grid(row=1, column=0, columnspan=3, sticky=tk.W, padx=20, pady=5)

lbl_res_class = ttk.Label(tab2, text="Class: -", font=("Arial", 11))
lbl_res_class.grid(row=2, column=0, columnspan=3, sticky=tk.W, padx=20, pady=5)

lbl_res_total = ttk.Label(tab2, text="Total Marks: -", font=("Arial", 11))
lbl_res_total.grid(row=3, column=0, columnspan=3, sticky=tk.W, padx=20, pady=5)

lbl_res_per = ttk.Label(tab2, text="Percentage: -", font=("Arial", 11))
lbl_res_per.grid(row=4, column=0, columnspan=3, sticky=tk.W, padx=20, pady=5)

lbl_res_status = ttk.Label(tab2, text="Result: -", font=("Arial", 11, "bold"))
lbl_res_status.grid(row=5, column=0, columnspan=3, sticky=tk.W, padx=20, pady=5)

# --- TAB 3: SHOW ALL RESULTS ---
tab3 = ttk.Frame(notebook)
notebook.add(tab3, text="3. Show All Results")

columns = ("Name", "Roll No", "Class", "Total", "Percentage", "Result")
tree = ttk.Treeview(tab3, columns=columns, show="headings")

for col in columns:
    tree.heading(col, text=col)
    tree.column(col, width=100, anchor=tk.CENTER)

tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

ttk.Button(tab3, text="📋 Refresh / Show All Results", command=show_all_results).pack(pady=10)

root.mainloop()